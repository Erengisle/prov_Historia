#!/usr/bin/env python3
"""Bygger elevprov och facit (Word, ev. PDF) från en prov.json.

Användning:
    python3 bygg_prov.py prov/det-langa-1800-talet-gy25/prov.json [--pdf] [--ut MAPP]
"""
import argparse
import json
import subprocess
import sys
from collections import Counter
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("python-docx saknas. Kör: pip install python-docx")

MIN_FLERVAL, MAX_FLERVAL = 12, 15
ANTAL_ALTERNATIV = 3
RUTA, KRYSS = "☐", "☒"  # ☐ ☒
BOKSTAVER = "ABCDEFGH"
GRA = RGBColor(0x55, 0x55, 0x55)


# ---------- Kontroll ----------

def kontrollera(prov):
    fel, varningar = [], []
    delar = prov.get("delar") or []
    if not prov.get("titel"):
        fel.append("'titel' saknas.")
    if not delar:
        fel.append("'delar' saknas eller är tom.")

    flerval = [f for d in delar if d.get("typ") == "flerval" for f in d.get("fragor", [])]
    if not MIN_FLERVAL <= len(flerval) <= MAX_FLERVAL:
        varningar.append(f"{len(flerval)} flervalsfrågor (önskat {MIN_FLERVAL}–{MAX_FLERVAL}).")
    for i, f in enumerate(flerval, 1):
        alt = f.get("alternativ", [])
        if len(alt) != ANTAL_ALTERNATIV:
            fel.append(f"Flervalsfråga {i}: {len(alt)} alternativ (ska vara {ANTAL_ALTERNATIV}).")
        if not isinstance(f.get("ratt"), int) or not 0 <= f["ratt"] < len(alt):
            fel.append(f"Flervalsfråga {i}: ogiltigt 'ratt' ({f.get('ratt')!r}).")
        if len(set(alt)) != len(alt):
            fel.append(f"Flervalsfråga {i}: två alternativ är identiska.")
        if alt and isinstance(f.get("ratt"), int) and 0 <= f["ratt"] < len(alt):
            if len(alt[f["ratt"]]) > 1.6 * max(len(a) for j, a in enumerate(alt) if j != f["ratt"]):
                varningar.append(f"Flervalsfråga {i}: rätt alternativ är mycket längre än de andra.")
    if flerval:
        fordelning = Counter(f.get("ratt") for f in flerval)
        if max(fordelning.values()) > len(flerval) * 0.5:
            varningar.append(f"Rätt svar är ojämnt fördelat på positionerna: {dict(sorted(fordelning.items()))}.")

    for n, d in enumerate(delar, 1):
        typ = d.get("typ")
        if typ == "begrepp":
            for b in d.get("begrepp", []):
                if not b.get("facit"):
                    varningar.append(f"Del {n}: begreppet '{b.get('term')}' saknar facit.")
        elif typ == "fritext":
            if not d.get("fraga"):
                fel.append(f"Del {n}: fritextfråga saknar 'fraga'.")
            if not d.get("bedomning"):
                varningar.append(f"Del {n}: fritextfråga saknar 'bedomning' (E/C/A).")
            if not d.get("kriterier"):
                varningar.append(f"Del {n}: fritextfråga saknar 'kriterier' (vilka betygskriterier den prövar).")
        elif typ != "flerval":
            fel.append(f"Del {n}: okänd typ {typ!r} (flerval, begrepp, fritext).")
    return fel, varningar


def poang(prov):
    summa = 0
    for d in prov["delar"]:
        if d["typ"] == "flerval":
            summa += sum(f.get("poang", 1) for f in d["fragor"])
        elif d["typ"] == "begrepp":
            summa += sum(b.get("poang", 2) for b in d["begrepp"])
        else:
            summa += d.get("poang", 0)
    return summa


# ---------- Hjälpfunktioner för Word ----------

def nytt_dokument(prov, etikett):
    doc = Document()
    sektion = doc.sections[0]
    sektion.page_height, sektion.page_width = Cm(29.7), Cm(21.0)
    sektion.left_margin = sektion.right_margin = Cm(2.2)
    sektion.top_margin, sektion.bottom_margin = Cm(1.8), Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(0)

    # Sidhuvud: etikett + kapitel till höger
    huvud = sektion.header.paragraphs[0]
    huvud.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = huvud.add_run(etikett)
    r.bold, r.font.size = True, Pt(12)
    if prov.get("rubrik_hoger"):
        r = huvud.add_run("\n" + prov["rubrik_hoger"])
        r.font.size, r.font.color.rgb = Pt(8), GRA

    # Sidfot: sidnummer
    fot = sektion.footer.paragraphs[0]
    fot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kod in ("begin", None, "end"):
        run = fot.add_run()
        if kod is None:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = "PAGE"
            run._r.append(instr)
        else:
            tecken = OxmlElement("w:fldChar")
            tecken.set(qn("w:fldCharType"), kod)
            run._r.append(tecken)
    return doc


def titel(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(26)


def rubrik(doc, text, fore=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(fore)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(12.5)
    return p


def stycke(doc, text="", fet=False, storlek=11, fore=0, efter=0, farg=None, hall_ihop=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(fore)
    p.paragraph_format.space_after = Pt(efter)
    p.paragraph_format.keep_with_next = hall_ihop
    if text:
        r = p.add_run(text)
        r.bold, r.font.size = fet, Pt(storlek)
        if farg:
            r.font.color.rgb = farg
    return p


def understruken(p, farg="808080"):
    """Lägger en linje under stycket (används som skrivrad)."""
    ppr = p._p.get_or_add_pPr()
    kant = OxmlElement("w:pBdr")
    under = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", farg)):
        under.set(qn(k), v)
    kant.append(under)
    ppr.append(kant)


def skrivrader(doc, antal, forsta_text=None):
    for i in range(antal):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = Pt(26)
        p.paragraph_format.space_after = Pt(0)
        if i == 0 and forsta_text:
            p.add_run(forsta_text).font.size = Pt(10.5)
        if i < antal - 1:
            p.paragraph_format.keep_with_next = True
        understruken(p)


def sidbrytning(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def namnrad(doc):
    tabell = doc.add_table(rows=1, cols=2)
    tabell.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(tabell.rows[0].cells, ("Namn:", "Klass:")):
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = Pt(24)
        p.add_run(text).font.size = Pt(10.5)
        understruken(p)
    tabell.rows[0].cells[0].width, tabell.rows[0].cells[1].width = Cm(11.5), Cm(5)
    stycke(doc, efter=4)


def flervalsfraga(doc, f, facit):
    p = stycke(doc, f["fraga"], storlek=11.5, fore=9, efter=3, hall_ihop=True)
    p.paragraph_format.keep_together = True
    for j, alt in enumerate(f["alternativ"]):
        ratt = facit and j == f["ratt"]
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.1)
        p.paragraph_format.keep_with_next = j < len(f["alternativ"]) - 1
        r = p.add_run(f"{KRYSS if ratt else RUTA}  {alt}")
        r.font.size, r.bold = Pt(10.5), ratt


def poangtext(p):
    return f"  ({p} p)" if p else ""


# ---------- Kriterier och centralt innehåll (numrerade, K1… och CI1…) ----------

SKILLMAPP = Path(__file__).resolve().parent.parent


def las_kriterier(prov):
    """Läser kriterier-gy11.json / kriterier-gy25.json utifrån provets utgåva."""
    utgava = (prov.get("utgava") or "").lower().replace("-", "")
    fil = SKILLMAPP / f"kriterier-{utgava}.json"
    return json.loads(fil.read_text(encoding="utf-8")) if utgava and fil.exists() else None


def med_namn(koder, lista):
    """'K2' -> 'K2 Samband …'. Text som inte är en känd kod skrivs som den är."""
    return [f"{k} {lista[k]}" if lista and k in lista else k for k in koder]


def oversikt(doc, prov, kr):
    """Tabeller över vilka kriterier och vilken del av det centrala innehållet varje del av provet prövar."""
    for rubriktext, nyckel in (("Betygskriterier som prövas", "kriterier"),
                               ("Centralt innehåll som prövas", "centralt_innehall")):
        lista = kr[nyckel]
        delar = {k: [str(n) for n, d in enumerate(prov["delar"], 1) if k in d.get(nyckel, [])] for k in lista}
        stycke(doc, f"{rubriktext} (kursplan {kr['kursplan']})", fet=True, fore=8, efter=2, hall_ihop=True)
        tabell = doc.add_table(rows=1, cols=3)
        tabell.style = "Table Grid"
        for cell, text in zip(tabell.rows[0].cells, ("Nr", "Innehåll", "Prövas i del")):
            cell.paragraphs[0].add_run(text).bold = True
        for k, namn in lista.items():
            rad = tabell.add_row().cells
            rad[0].paragraphs[0].add_run(k).bold = True
            rad[1].paragraphs[0].add_run(namn).font.size = Pt(10)
            rad[2].paragraphs[0].add_run(", ".join(delar[k]) or "–").font.size = Pt(10)
            if not delar[k]:
                for c in rad:
                    for r in c.paragraphs[0].runs:
                        r.font.color.rgb = GRA
        for rad in tabell.rows:
            rad.cells[0].width, rad.cells[1].width, rad.cells[2].width = Cm(1.4), Cm(12.4), Cm(2.8)


# ---------- Elevversion ----------

def bygg_elev(prov, sokvag):
    doc = nytt_dokument(prov, "ELEV")
    titel(doc, prov["titel"])
    namnrad(doc)
    for n, d in enumerate(prov["delar"], 1):
        typ = d["typ"]
        if typ == "flerval":
            rubrik(doc, f"{n}. {d.get('instruktion', 'Välj rätt alternativ.')}", fore=4)
            for f in d["fragor"]:
                flervalsfraga(doc, f, facit=False)
        elif typ == "begrepp":
            sidbrytning(doc)
            rubrik(doc, f"{n}. {d.get('instruktion', 'Förklara kortfattat med egna ord.')}", fore=0)
            for b in d["begrepp"]:
                stycke(doc, efter=4)
                skrivrader(doc, d.get("rader", 3), forsta_text=b["term"])
        elif typ == "fritext":
            rader = d.get("rader", 8)
            if rader >= 15:
                sidbrytning(doc)
            p = rubrik(doc, f"{n}. {d['fraga']}", fore=0 if rader >= 15 else 22)
            stycke(doc, efter=6, hall_ihop=True)
            skrivrader(doc, rader)
    doc.save(sokvag)


# ---------- Facit ----------

def bygg_facit(prov, sokvag):
    doc = nytt_dokument(prov, "FACIT")
    titel(doc, f"{prov['titel']} – facit")
    info = f"Totalt {poang(prov)} poäng."
    if prov.get("kalla"):
        info = f"Underlag: {prov['kalla']}. " + info
    if prov.get("kursplan"):
        info += f" Bedömningsstöd enligt kursplan {prov['kursplan']}."
    stycke(doc, info, storlek=10, efter=6, farg=GRA)
    stycke(doc, "Rättningsnyckel flerval (numrering som i Trelson-filen): " + rattningsnyckel(prov),
           storlek=10, efter=6)
    kr = las_kriterier(prov)
    if kr:
        oversikt(doc, prov, kr)

    def kodrader(d):
        for etikett, nyckel in (("Betygskriterier", "kriterier"), ("Centralt innehåll", "centralt_innehall")):
            if d.get(nyckel):
                stycke(doc, f"{etikett}: " + "; ".join(med_namn(d[nyckel], kr and kr[nyckel])),
                       storlek=10, efter=2, farg=GRA)

    for n, d in enumerate(prov["delar"], 1):
        typ = d["typ"]
        if typ == "flerval":
            antal = sum(f.get("poang", 1) for f in d["fragor"])
            rubrik(doc, f"{n}. {d.get('instruktion', 'Välj rätt alternativ.')}{poangtext(antal)}")
            kodrader(d)
            for f in d["fragor"]:
                flervalsfraga(doc, f, facit=True)
        elif typ == "begrepp":
            antal = sum(b.get("poang", 2) for b in d["begrepp"])
            rubrik(doc, f"{n}. {d.get('instruktion', 'Förklara kortfattat med egna ord.')}{poangtext(antal)}")
            kodrader(d)
            for b in d["begrepp"]:
                stycke(doc, b["term"] + poangtext(b.get("poang", 2)), fet=True, fore=6, hall_ihop=True)
                stycke(doc, b.get("facit", ""), storlek=10.5)
        elif typ == "fritext":
            rubrik(doc, f"{n}. {d['fraga']}{poangtext(d.get('poang'))}")
            if d.get("historiska_begrepp"):
                stycke(doc, "Historiska begrepp: " + ", ".join(d["historiska_begrepp"]), storlek=10, efter=2, farg=GRA)
            kodrader(d)
            if d.get("fordjupning"):
                stycke(doc, f"Bygger på fördjupningen {d['fordjupning']}. Frågan ger själv den bakgrund som behövs.",
                       storlek=10.5, efter=2, farg=GRA)
            if d.get("innehall"):
                stycke(doc, "Ett bra svar tar upp:", fet=True, hall_ihop=True)
                for punkt in d["innehall"]:
                    p = doc.add_paragraph(style="List Bullet")
                    p.add_run(punkt).font.size = Pt(10.5)
            if d.get("bedomning"):
                stycke(doc, "Bedömningsstöd", fet=True, fore=6, efter=2, hall_ihop=True)
                tabell = doc.add_table(rows=0, cols=2)
                tabell.style = "Table Grid"
                for niva in ("E", "C", "A"):
                    if niva in d["bedomning"]:
                        rad = tabell.add_row().cells
                        rad[0].width, rad[1].width = Cm(1.2), Cm(15.4)
                        rad[0].paragraphs[0].add_run(niva).bold = True
                        rad[1].paragraphs[0].add_run(d["bedomning"][niva]).font.size = Pt(10.5)
    doc.save(sokvag)



# ---------- Trelson (en fråga i taget, för att kopiera in i provverktyget) ----------

def trelsonfragor(prov):
    """Provet som en platt lista av enskilda frågor med löpande nummer.
    Numreringen används i Trelson-filen, i rättningsnyckeln och vid rättning."""
    nr = 0
    for d in prov["delar"]:
        if d["typ"] == "flerval":
            for f in d["fragor"]:
                nr += 1
                yield {"nr": nr, "typ": "Flerval", "poang": f.get("poang", 1), "text": f["fraga"],
                       "alternativ": f["alternativ"], "ratt": f["ratt"]}
        elif d["typ"] == "begrepp":
            instruktion = d.get("instruktion", "Förklara kortfattat med egna ord.").rstrip(".:")
            for b in d["begrepp"]:
                nr += 1
                yield {"nr": nr, "typ": "Kort svar", "poang": b.get("poang", 2),
                       "text": f"{instruktion}: {b['term']}"}
        else:
            nr += 1
            yield {"nr": nr, "typ": "Längre svar" if d.get("rader", 8) >= 15 else "Svar med några meningar",
                   "poang": d.get("poang", 0), "text": d["fraga"]}


def bygg_trelson(prov, docx_sokvag, txt_sokvag):
    rader = [prov["titel"], ""]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph().add_run(prov["titel"]).bold = True
    stycke(doc, "En fråga per block. Kopiera frågetexten och alternativen var för sig in i Trelson.",
           storlek=9.5, efter=8, farg=GRA)
    for f in trelsonfragor(prov):
        huvud = f"Fråga {f['nr']} – {f['typ']} ({f['poang']} p)"
        rader += [huvud, f["text"]]
        stycke(doc, huvud, fet=True, storlek=10, fore=10, farg=GRA, hall_ihop=True)
        stycke(doc, f["text"], hall_ihop="alternativ" in f)
        for j, alt in enumerate(f.get("alternativ", [])):
            rader.append(f"{BOKSTAVER[j]}. {alt}")
            stycke(doc, f"{BOKSTAVER[j]}. {alt}", hall_ihop=j < len(f["alternativ"]) - 1)
        rader.append("")
    doc.save(docx_sokvag)
    txt_sokvag.write_text("\n".join(rader), encoding="utf-8")


def rattningsnyckel(prov):
    return "   ".join(f"{f['nr']} {BOKSTAVER[f['ratt']]}" for f in trelsonfragor(prov) if "alternativ" in f)


# ---------- Huvudprogram ----------

def till_pdf(docx, mapp):
    pdf = mapp / (docx.stem + ".pdf")
    pdf.unlink(missing_ok=True)
    try:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(mapp), str(docx)],
                       capture_output=True, timeout=180)
    except (OSError, subprocess.SubprocessError) as e:
        print(f"VARNING: kunde inte göra PDF av {docx.name}: {e}", file=sys.stderr)
        return
    if not pdf.exists():  # soffice avslutar med 0 även när konverteringen misslyckas
        print(f"VARNING: LibreOffice gjorde ingen PDF av {docx.name}. Spara som PDF från Word i stället.",
              file=sys.stderr)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json", type=Path)
    ap.add_argument("--ut", type=Path, help="Utmapp (standard: samma mapp som JSON-filen)")
    ap.add_argument("--pdf", action="store_true", help="Gör även PDF (kräver LibreOffice)")
    args = ap.parse_args()

    prov = json.loads(args.json.read_text(encoding="utf-8"))
    fel, varningar = kontrollera(prov)
    for v in varningar:
        print("VARNING:", v)
    if fel:
        for f in fel:
            print("FEL:", f)
        sys.exit(1)

    mapp = args.ut or args.json.parent
    mapp.mkdir(parents=True, exist_ok=True)
    namn = args.json.parent.name if args.json.stem == "prov" else args.json.stem
    elev, facit = mapp / f"{namn}_elev.docx", mapp / f"{namn}_facit.docx"
    bygg_elev(prov, elev)
    bygg_facit(prov, facit)
    bygg_trelson(prov, mapp / f"{namn}_trelson.docx", mapp / f"{namn}_trelson.txt")
    if args.pdf:
        till_pdf(elev, mapp)
        till_pdf(facit, mapp)

    antal_fv = sum(len(d["fragor"]) for d in prov["delar"] if d["typ"] == "flerval")
    print(f"Klart: {antal_fv} flervalsfrågor, {len(prov['delar'])} delar, {poang(prov)} poäng.")
    for p in sorted(mapp.glob(f"{namn}_*")):
        print("  ", p)


if __name__ == "__main__":
    main()
